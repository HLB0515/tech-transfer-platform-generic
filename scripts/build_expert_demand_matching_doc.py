from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = "/Users/macbookpro/Documents/New project/docs/企业需求与专家匹配评分体系说明.docx"


def set_cell_shading(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(text))
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(9.5)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_para(doc, text, style=None, bold=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.1
    run = p.add_run(text)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.bold = bold
    run.font.size = Pt(11)
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_text(hdr[i], h, bold=True, color="FFFFFF")
        set_cell_shading(hdr[i], "1F4D78")
        if widths:
            hdr[i].width = Inches(widths[i])
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
            if widths:
                cells[i].width = Inches(widths[i])
    doc.add_paragraph()
    return table


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)

styles = doc.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
styles["Normal"].font.size = Pt(11)
for name, size, color in [("Heading 1", 16, "2E74B5"), ("Heading 2", 13, "2E74B5"), ("Heading 3", 12, "1F4D78")]:
    styles[name].font.name = "Calibri"
    styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    styles[name].font.size = Pt(size)
    styles[name].font.color.rgb = RGBColor.from_string(color)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("企业需求与专家匹配评分体系说明")
run.font.name = "Calibri"
run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
run.font.size = Pt(20)
run.bold = True
run.font.color.rgb = RGBColor.from_string("0B2545")

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub_run = sub.add_run("适用于科技成果转化平台需求库与专家库的一对一智能匹配、分层推荐与技术经理人介入判断")
sub_run.font.name = "Calibri"
sub_run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
sub_run.font.size = Pt(10.5)
sub_run.font.color.rgb = RGBColor.from_string("4B6380")

doc.add_heading("一、建设目标", level=1)
for text in [
    "本体系用于在企业技术需求通过真实性评分后，进一步判断其与专家库中专家能力画像的匹配程度，形成可解释、可复核、可迭代的专家推荐结果。",
    "平台不采用简单关键词命中作为唯一依据，而是综合技术领域、语义相似度、技术指标对应关系、专家研究成果、工程化经验、产业链适配度、需求真实性和复杂度等因素进行评分。",
    "最终输出0—100分匹配分数、等级标签、推荐理由、风险提示和是否建议技术经理人介入。"
]:
    add_para(doc, text)

doc.add_heading("二、数据基础", level=1)
add_table(
    doc,
    ["数据对象", "主要字段", "用于匹配的核心信息"],
    [
        ["企业需求库", "需求名称、企业名称、技术领域、需求详情、技术指标、技术难题、现有基础、预算、地区、有效期", "识别企业真实痛点、技术指标、攻关边界、实施基础和预算投入意愿"],
        ["专家库", "专家名称、专业特长、战略性新兴产业分类、研究方向、所属单位、职务职称、工作业绩、相关研究成果", "形成专家能力画像，判断其是否具备解决该类技术问题的专业基础和成果支撑"],
        ["辅助数据", "成果库、政策库、产业链词库、专家评分样本、历史撮合记录", "用于训练语义模型、校准评分权重、补充产业链和成果转化判断"],
    ],
    widths=[1.2, 2.7, 2.6],
)

doc.add_heading("三、匹配流程", level=1)
for text in [
    "第一步：真实性前置筛查。企业需求真实性低于60分的，不进入自动专家推荐，先退回补充、人工复核或由技术经理人访谈澄清。",
    "第二步：专家候选召回。根据技术领域、产业分类、关键词、研究方向和专业特长，先从专家库中召回一批可能相关专家。",
    "第三步：多维度精排打分。对候选专家计算综合匹配分，并生成推荐理由和风险标签。",
    "第四步：复杂项目路由。匹配分较高但技术指标复杂、预算较高、涉及中试/知识产权/质量责任的项目，建议技术经理人介入撮合。",
    "第五步：结果反馈学习。将专家确认、企业反馈、磋商结果、签约结果回流，用于持续优化权重和语义模型。"
]:
    add_para(doc, text)

doc.add_heading("四、百分制评分维度", level=1)
add_table(
    doc,
    ["评分维度", "分值", "评价重点", "高分表现", "低分风险"],
    [
        ["技术领域与产业链适配", "20分", "需求技术领域与专家产业分类、研究方向、专业类别是否一致或相邻", "同属先进制造、新材料、生物医药等方向，且处在同一产业链环节", "领域跨度过大，仅有泛泛相关"],
        ["语义相似度与专业概念匹配", "20分", "通过AI语义模型识别专业词、同义词、上下游概念和隐含技术关系", "即使字面不完全相同，也能识别材料、工艺、装备、算法等专业关联", "仅有少量普通词重合"],
        ["技术指标与专家能力对应", "20分", "需求中的性能指标、工艺指标、检测指标是否能对应专家专长和成果", "专家研究方向可解释关键指标，如强度、耐腐蚀、良率、准确率、寿命等", "专家方向宽泛，无法对应具体指标"],
        ["专家成果与项目支撑", "15分", "专家是否有论文、专利、项目、平台、工程案例或成果转化经历支撑", "有相关课题、成果、专利、实验平台或中试经验", "只有学科方向相近，缺少可验证成果"],
        ["工程化与转化适配", "10分", "专家能力是否能支撑样件、中试、量产、检测、标准、认证等转化环节", "具备工程化、产业化或联合研发经验", "偏基础研究，短期难以对接企业场景"],
        ["地域与协作可达性", "5分", "专家所在地区、合作网络和长三角协作便利性", "具备长三角或线上线下协作条件", "协作成本较高但不作为硬性排除"],
        ["需求真实性与复杂度校正", "10分", "结合需求真实性评分和技术复杂度进行校正", "真实需求、指标清楚、复杂度高且需要专家深度介入", "需求真实性不足或需求本身更适合成熟采购"],
    ],
    widths=[1.35, 0.55, 1.55, 1.55, 1.5],
)

doc.add_heading("五、分数等级与处置规则", level=1)
add_table(
    doc,
    ["等级", "分数区间", "含义", "平台处置"],
    [
        ["S级", "90—100分", "超级高匹配。专家方向、成果、指标和转化能力高度一致。", "列入重点撮合，优先发起双方智能体磋商，并建议形成转化撮合报告。"],
        ["A级", "80—89分", "匹配很高。具备明确专业相关性和较强成果支撑。", "优先推荐给企业，可进入一对一预沟通，必要时拉入技术经理人。"],
        ["B级", "70—79分", "匹配较高。方向相关，但部分指标、成果或工程化能力需确认。", "推荐进入第一轮预沟通，由智能体补充追问。"],
        ["C级", "60—69分", "匹配凑合。可能存在相邻领域专家或间接相关专家。", "作为备选专家，不宜直接作为首推，建议补充访谈后再判断。"],
        ["D级", "0—59分", "匹配不足。领域、语义或能力支撑较弱。", "不进入自动推荐，除非人工指定或作为跨界咨询专家。"],
    ],
    widths=[0.7, 0.9, 2.2, 2.7],
)

doc.add_heading("六、技术经理人介入规则", level=1)
add_para(doc, "专家匹配并不等于项目可以直接成交。以下情形即使匹配分较高，也建议由技术经理人介入，负责澄清边界、组织磋商、控制风险。")
add_table(
    doc,
    ["触发条件", "判断依据", "技术经理人任务"],
    [
        ["复杂研发项目", "技术指标多、涉及中试、样件、批量稳定性、质量责任", "拆解攻关任务，形成里程碑、验收指标和样件验证计划"],
        ["高预算或股权/长期合作", "预算较高、合作方式涉及股权投资、长期供应或联合开发", "协助谈判合作模式、报价、投入边界和责任分配"],
        ["知识产权或合规风险", "涉及专利、认证、标准、临床、军工、适航等", "组织专利检索、合规审查和保密协议"],
        ["专家匹配高但企业需求不够清晰", "匹配分高，但需求真实性仅60—70分或指标不完整", "访谈企业，补充技术指标和真实场景后再正式对接"],
    ],
    widths=[1.5, 2.4, 2.6],
)

doc.add_heading("七、算法实现建议", level=1)
add_table(
    doc,
    ["模块", "方法", "说明"],
    [
        ["规则筛选", "字段完整性、领域一致性、硬门槛过滤", "快速排除低真实性需求和明显不相关专家"],
        ["关键词模型", "TF-IDF、BM25、专业词典", "识别直接词面相关，如农业大棚塑料—高强度棚膜"],
        ["语义模型", "中文向量模型/大模型Embedding", "识别字面不同但专业相关的需求，如铝合金强韧化—微合金化材料专家"],
        ["AI复核", "大模型重排序与理由生成", "对Top候选专家生成推荐理由、风险提示和下一轮追问"],
        ["人工校准", "专家评分、历史撮合结果、技术经理人反馈", "持续调整权重，提升科学性和可信度"],
    ],
    widths=[1.2, 2.1, 3.2],
)

doc.add_heading("八、评分公式示意", level=1)
add_para(doc, "专家匹配综合分 = 领域产业链适配20% + 语义相似度20% + 技术指标对应20% + 专家成果支撑15% + 工程化转化适配10% + 地域协作5% + 需求真实性与复杂度校正10%。")
add_para(doc, "其中，需求真实性低于60分时不建议进入自动匹配；真实性为60—79分时可以匹配但需补充材料；真实性80分以上可进入优先推荐。")

doc.add_heading("九、输出结果样式", level=1)
add_table(
    doc,
    ["输出项", "示例"],
    [
        ["匹配分数", "86分"],
        ["等级", "A级：匹配很高"],
        ["推荐理由", "专家研究方向覆盖高强铝合金、材料组织调控和耐腐蚀性能，与企业需求中的抗拉强度、盐雾试验和焊接性能指标高度对应。"],
        ["风险提示", "仍需确认专家是否具备中试放大和结构件焊接工艺经验。"],
        ["推荐动作", "建议发起专家智能体预沟通，并拉入技术经理人确认样件、NDA、报价和里程碑。"],
    ],
    widths=[1.3, 5.2],
)

doc.add_heading("十、对外汇报表述", level=1)
add_para(doc, "该匹配体系的核心不是把专家库简单搜索一遍，而是把企业需求转化为结构化技术画像，再与专家的研究方向、专业特长、成果支撑和工程化能力进行多维度比对。平台既能识别高强相关的直接匹配，也能通过AI语义模型发现字面不相同但技术上高度相关的潜在专家。对于复杂项目，系统不会简单推送专家，而是建议技术经理人介入，形成“需求识别—专家匹配—智能体磋商—经理人撮合—转化报告”的闭环。")

doc.save(OUT)
print(OUT)
