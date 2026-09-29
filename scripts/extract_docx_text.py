from pathlib import Path

from docx import Document


def main():
    docx_path = Path("/Users/macbookpro/Desktop/科技成果转化平台/智能体广场设计说明.docx")
    doc = Document(docx_path)
    lines = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            lines.append(text)
    for index, table in enumerate(doc.tables, start=1):
        lines.append(f"【表格{index}】")
        for row in table.rows:
            values = [cell.text.strip().replace("\n", " / ") for cell in row.cells]
            if any(values):
                lines.append(" | ".join(values))
    print("\n".join(lines))


if __name__ == "__main__":
    main()
