import json
import re
import shutil
from datetime import datetime, date
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path("/Users/macbookpro/Documents/New project")
STORE_PATH = ROOT / "data" / "store.json"
DEMAND_PATH = Path("/Users/macbookpro/Desktop/科技成果转化平台/AAA/（需求库）.xlsx")
EXPERT_PATH = Path("/Users/macbookpro/Desktop/（专家库）清华专家信息表-相关研究成果增强版.xlsx")


def text(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def read_rows(path, sheet_index=0):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[sheet_index]
    rows = []
    for row in ws.iter_rows(values_only=True):
        if any(cell not in (None, "") for cell in row):
            rows.append(list(row))
    return rows


def parse_money(value):
    if value is None or value == "":
        return 0
    if isinstance(value, (int, float)):
        # 当前需求库预算以“万元”为口径，网站内部统一转成“元”。
        return int(float(value) * 10000)
    raw = str(value).replace(",", "").strip()
    nums = re.findall(r"\d+(?:\.\d+)?", raw)
    if not nums:
        return 0
    amount = float(nums[0])
    if "亿" in raw:
        return int(amount * 100000000)
    if "万" in raw:
        return int(amount * 10000)
    return int(amount)


def authenticity(demand):
    joined = "".join(
        text(demand.get(k))
        for k in [
            "name",
            "company",
            "field",
            "detail",
            "technicalIndicators",
            "problem",
            "foundation",
            "cooperation",
            "orgType",
            "budgetText",
            "demandType",
            "region",
        ]
    )
    vague = re.compile(r"未说明|面议|暂无|待定|不详|根据需要|看情况|越高越好|越低越好|先进水平")
    score = 0
    risk = []
    strengths = []

    core = [demand.get("name"), demand.get("detail"), demand.get("technicalIndicators"), demand.get("problem"), demand.get("foundation")]
    filled = sum(1 for item in core if text(item))
    score += min(20, filled * 4)
    if filled >= 4:
        strengths.append("关键字段较完整")
    else:
        risk.append("关键字段缺失")

    indicators = text(demand.get("technicalIndicators"))
    nums = re.findall(r"\d+(?:\.\d+)?|≥|≤|%|MPa|GPa|ppm|℃|小时|天|吨|m2|mm|μm|微米", indicators)
    if len(nums) >= 6:
        score += 20
        strengths.append("技术指标充分")
    elif len(nums) >= 3:
        score += 15
        strengths.append("有可验证技术指标")
    elif len(nums) >= 1:
        score += 8
    else:
        risk.append("技术指标不可验证")

    problem = text(demand.get("problem"))
    detail = text(demand.get("detail"))
    if len(problem) >= 80 or len(detail) >= 160:
        score += 15
        strengths.append("技术问题描述较清楚")
    elif len(problem) >= 30 or len(detail) >= 80:
        score += 10
    elif len(problem) >= 10:
        score += 5
    else:
        risk.append("技术问题空泛")

    foundation = text(demand.get("foundation"))
    if re.search(r"员工|研发|设备|产线|车间|实验室|检测|中试|专利|团队|客户|订单|样品|样件", foundation):
        score += 15
        strengths.append("现有基础较扎实")
    elif foundation and not vague.search(foundation):
        score += 7
    else:
        risk.append("现有基础不足")

    budget_amount = int(demand.get("budgetAmount") or 0)
    if budget_amount >= 1000000:
        score += 10
        strengths.append("预算投入明确")
    elif budget_amount > 0:
        score += 6
    else:
        risk.append("预算金额未量化")

    tech_need = f"{demand.get('name','')} {demand.get('detail','')} {demand.get('technicalIndicators','')} {demand.get('problem','')}"
    if re.search(r"中试|产业化|批量|稳定性|可靠性|国产化|卡脖子|专利|标准|新材料|新工艺|新装备|算法|检测|认证", tech_need):
        score += 10
        strengths.append("具备攻关或转化必要性")
    elif re.search(r"采购|购买|设备", tech_need) and not re.search(r"研发|开发|攻关|优化", tech_need):
        score += 3
        risk.append("成熟技术可替代")
    else:
        score += 5

    consistency = 10
    if vague.search(joined):
        consistency -= 3
        risk.append("存在模糊表述")
    if len(joined) < 80:
        consistency -= 4
        risk.append("文本信息量不足")
    score += max(0, consistency)

    score = max(0, min(100, round(score)))
    if score >= 80:
        level, recommendation = "高可信", "优先对接"
    elif score >= 60:
        level, recommendation = "基本可信", "补充访谈后对接"
    elif score >= 40:
        level, recommendation = "疑似虚假", "人工复核并要求补证"
    else:
        level, recommendation = "虚假/高风险", "判定高风险，暂缓入库"
    return {
        "score": score,
        "level": level,
        "isReal": score >= 60,
        "verdict": "真实需求" if score >= 60 else "虚假需求",
        "recommendation": recommendation,
        "riskTags": sorted(set(risk)),
        "strengths": sorted(set(strengths)),
        "evaluatedAt": datetime.now().isoformat(timespec="seconds"),
    }


def import_demands():
    rows = read_rows(DEMAND_PATH)
    headers = [text(item) for item in rows[0]]
    index = {name: i for i, name in enumerate(headers)}

    def get(row, name):
        return text(row[index[name]]) if name in index and index[name] < len(row) else ""

    demands = []
    for row in rows[1:]:
        if not get(row, "需求名称"):
            continue
        budget_raw = row[index["预算金额"]] if "预算金额" in index and index["预算金额"] < len(row) else ""
        budget_amount = parse_money(budget_raw)
        valid = get(row, "有效期")
        updated = get(row, "更新时间")
        field = get(row, "技术领域（先进制造业、农业、生物医药、现代服务业及其他）") or get(row, "技术领域")
        item = {
            "id": len(demands) + 1,
            "name": get(row, "需求名称"),
            "company": get(row, "所属单位或公司名称"),
            "field": field,
            "detail": get(row, "需求详情"),
            "technicalIndicators": get(row, "技术指标"),
            "problem": get(row, "拟解决的技术难题"),
            "foundation": get(row, "现有基础条件"),
            "budget": budget_amount,
            "cooperation": get(row, "合作方式"),
            "orgType": get(row, "单位性质"),
            "budgetAmount": budget_amount,
            "budgetText": f"{text(budget_raw)}万元" if isinstance(budget_raw, (int, float)) and budget_raw else text(budget_raw) or "面议",
            "demandType": get(row, "需求类型"),
            "region": get(row, "所属地区"),
            "validUntil": valid,
            "validUntilText": valid,
            "updatedAt": updated or datetime.now().isoformat(timespec="seconds"),
            "source": "（需求库）.xlsx",
            "keywords": ";".join(filter(None, [field, get(row, "需求类型"), get(row, "所属地区")])),
            "publisherId": 1,
        }
        item["authenticity"] = authenticity(item)
        demands.append(item)
    return demands


def import_experts():
    rows = read_rows(EXPERT_PATH, 0)
    headers = [text(item) or "专家介绍" for item in rows[0]]
    index = {name: i for i, name in enumerate(headers)}

    def get(row, name):
        return text(row[index[name]]) if name in index and index[name] < len(row) else ""

    experts = []
    for row in rows[1:]:
        name = get(row, "专家名称")
        if not name:
            continue
        related = get(row, "相关研究成果")
        achievements = get(row, "工作业绩及专业资质")
        if related:
            achievements = f"{achievements}\n相关研究成果：{related}".strip()
        bio = get(row, "个人简介")
        if related and related not in bio:
            bio = f"{bio}\n相关研究成果：{related}".strip()
        research = get(row, "研究方向")
        specialties = get(row, "专业特长")
        item = {
            "id": len(experts) + 1,
            "intro": get(row, "专家介绍"),
            "name": name,
            "keywords": specialties,
            "field": get(row, "战略性新兴产业分类"),
            "achievements": achievements,
            "region": get(row, "所属地域") or get(row, "所属地域2"),
            "researchDirection": research,
            "organization": get(row, "所属单位"),
            "orgType": get(row, "单位性质"),
            "position": get(row, "职务"),
            "department": get(row, "现所在部门"),
            "category": get(row, "专业类别"),
            "title": get(row, "专业技术职称"),
            "gender": get(row, "性别"),
            "education": get(row, "最高学历"),
            "graduatedFrom": get(row, "毕业院校"),
            "contact": get(row, "联系方式"),
            "bio": bio,
            "relatedResults": related,
            "activeScore": 92 if "院士" in f"{get(row, '职务')} {get(row, '专业技术职称')} {achievements}" else 86,
            "source": "（专家库）清华专家信息表-相关研究成果增强版.xlsx",
        }
        experts.append(item)
    return experts


def main():
    if not DEMAND_PATH.exists():
        raise FileNotFoundError(f"找不到需求库：{DEMAND_PATH}")
    if not EXPERT_PATH.exists():
        raise FileNotFoundError(f"找不到专家库：{EXPERT_PATH}")

    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup = ROOT / "data" / f"store.backup-before-latest-demand-expert-{timestamp}.json"
    shutil.copy2(STORE_PATH, backup)

    with STORE_PATH.open("r", encoding="utf-8") as f:
        store = json.load(f)

    demands = import_demands()
    experts = import_experts()
    store["demands"] = demands
    store["experts"] = experts
    store["matches"] = []
    # Keep this aligned with server.js SEED_VERSION so the server does not
    # append bundled demo records on the next startup.
    store["seedVersion"] = 8
    store["lastLibraryImport"] = {
        "demands": str(DEMAND_PATH),
        "experts": str(EXPERT_PATH),
        "demandCount": len(demands),
        "expertCount": len(experts),
        "importedAt": datetime.now().isoformat(timespec="seconds"),
        "backup": str(backup),
    }

    with STORE_PATH.open("w", encoding="utf-8") as f:
        json.dump(store, f, ensure_ascii=False, indent=2)
        f.write("\n")

    print(json.dumps(store["lastLibraryImport"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
