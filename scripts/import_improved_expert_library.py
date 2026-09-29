import json
import re
import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path("/Users/macbookpro/Documents/New project")
STORE_PATH = ROOT / "data" / "store.json"
EXPERT_PATH = Path("/Users/macbookpro/Desktop/（专家库）清华专家信息表-相关研究成果改进版.xlsx")


def text(value):
    if value is None:
        return ""
    return str(value).strip()


def clean_profile_text(value):
    value = text(value)
    value = re.sub(r"https?://\S+", "", value)
    value = re.sub(r"来源[:：]\s*[^；。\n]+[；。\n]?", "", value)
    value = re.sub(r"\s+", " ", value)
    return value.strip(" ；;，,")


def read_rows(path, sheet_index=0):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[sheet_index]
    rows = []
    for row in ws.iter_rows(values_only=True):
        if any(cell not in (None, "") for cell in row):
            rows.append(list(row))
    return rows


def import_experts():
    rows = read_rows(EXPERT_PATH, 0)
    headers = [text(item) or "专家介绍" for item in rows[0]]
    index = {name: i for i, name in enumerate(headers)}

    def get(row, name):
        return text(row[index[name]]) if name in index and index[name] < len(row) else ""

    experts = []
    for row in rows[1:]:
        name = get(row, "专家名称")
        if not name:
            continue
        related = clean_profile_text(get(row, "相关研究成果"))
        achievements = clean_profile_text(get(row, "工作业绩及专业资质"))
        research = clean_profile_text(get(row, "研究方向"))
        specialties = clean_profile_text(get(row, "专业特长"))
        bio = clean_profile_text(get(row, "个人简介"))

        # Keep the improved spreadsheet fields clean: no old URLs/source traces,
        # and no duplicated "来源：" text in the public expert profile.
        item = {
            "id": len(experts) + 1,
            "intro": clean_profile_text(get(row, "专家介绍")),
            "name": name,
            "keywords": specialties,
            "field": clean_profile_text(get(row, "战略性新兴产业分类")),
            "achievements": achievements,
            "region": clean_profile_text(get(row, "所属地域") or get(row, "所属地域2")),
            "researchDirection": research,
            "organization": clean_profile_text(get(row, "所属单位")),
            "orgType": clean_profile_text(get(row, "单位性质")),
            "position": clean_profile_text(get(row, "职务")),
            "department": clean_profile_text(get(row, "现所在部门")),
            "category": clean_profile_text(get(row, "专业类别")),
            "title": clean_profile_text(get(row, "专业技术职称")),
            "gender": clean_profile_text(get(row, "性别")),
            "education": clean_profile_text(get(row, "最高学历")),
            "graduatedFrom": clean_profile_text(get(row, "毕业院校")),
            "contact": clean_profile_text(get(row, "联系方式")),
            "bio": bio,
            "relatedResults": related,
            "activeScore": 92 if "院士" in f"{get(row, '职务')} {get(row, '专业技术职称')} {achievements}" else 86,
            "source": "清华大学专家库（相关研究成果改进版）",
        }
        experts.append(item)
    return experts


def main():
    if not EXPERT_PATH.exists():
        raise FileNotFoundError(f"找不到专家库：{EXPERT_PATH}")

    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup = ROOT / "data" / f"store.backup-before-improved-experts-{timestamp}.json"
    shutil.copy2(STORE_PATH, backup)

    with STORE_PATH.open("r", encoding="utf-8") as f:
        store = json.load(f)

    experts = import_experts()
    store["experts"] = experts
    store["matches"] = []
    store["seedVersion"] = 8
    store["lastExpertImport"] = {
        "experts": str(EXPERT_PATH),
        "expertCount": len(experts),
        "importedAt": datetime.now().isoformat(timespec="seconds"),
        "backup": str(backup),
    }

    with STORE_PATH.open("w", encoding="utf-8") as f:
        json.dump(store, f, ensure_ascii=False, indent=2)
        f.write("\n")

    print(json.dumps(store["lastExpertImport"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
